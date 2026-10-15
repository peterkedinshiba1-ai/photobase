"""PhotoBase - importe/exporte une base Excel (.xlsx) avec images intégrées
(Insertion > Images > "Placer sur les cellules") et les affiche en galerie."""
import base64
import io
import os
from pathlib import Path

import flet as ft
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image, ImageOps

DATA = Path(os.getenv("FLET_APP_STORAGE_DATA", "data"))
DATA.mkdir(parents=True, exist_ok=True)
BASE = DATA / "base.xlsx"  # copie de travail (dossier privé de l'appli)

FIELDS = ["nom", "tags", "album", "note"]
ALIASES = {
    "nom": ("nom", "name", "titre", "title"),
    "tags": ("tags", "tag", "mots-clés", "mots cles"),
    "album": ("album", "dossier", "catégorie", "categorie"),
    "note": ("note", "notes", "description", "commentaire"),
}
MAX_SIDE = 1600
Image.MAX_IMAGE_PIXELS = 60_000_000


def safe(v) -> str:
    s = str(v or "")
    return " " + s if s[:1] in ("=", "+", "-", "@") else s  # anti injection de formule


def normalize(raw: bytes):
    """Valide l'image, retire EXIF/GPS, renvoie (image_1600px, miniature)."""
    img = Image.open(io.BytesIO(raw))
    img.load()
    img = ImageOps.exif_transpose(img).convert("RGB")
    img.thumbnail((MAX_SIDE, MAX_SIDE))
    full = io.BytesIO()
    img.save(full, "JPEG", quality=85)
    t = img.copy()
    t.thumbnail((320, 320))
    tb = io.BytesIO()
    t.save(tb, "JPEG", quality=80)
    return full.getvalue(), tb.getvalue()


def new_item(nom="", tags="", album="", note="", raw=None):
    full, thumb = normalize(raw)
    return {"nom": nom, "tags": tags, "album": album, "note": note, "full": full, "thumb": thumb}


def read_xlsx(path) -> list[dict]:
    wb = load_workbook(path)
    out = []
    for ws in wb.worksheets:
        cols = {}
        for c in ws[1]:
            key = str(c.value or "").strip().lower()
            for f, names in ALIASES.items():
                if key in names:
                    cols[f] = c.column
        for im in ws._images:
            frm = getattr(im.anchor, "_from", None)
            if frm is None:
                continue
            row = frm.row + 1
            try:
                it = new_item(raw=im._data())
            except Exception:
                continue
            for f, col in cols.items():
                v = ws.cell(row=row, column=col).value
                it[f] = "" if v is None else str(v)[:500]
            if not it["nom"]:
                it["nom"] = f"{ws.title}-L{row}"
            out.append(it)
    return out


def write_xlsx(items: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "photos"
    ws.append([f for f in FIELDS] + ["photo"])
    for col, w in zip("ABCDE", (22, 22, 18, 30, 26)):
        ws.column_dimensions[col].width = w
    for i, it in enumerate(items, start=2):
        ws.append([safe(it[f]) for f in FIELDS])
        w, h = Image.open(io.BytesIO(it["full"])).size
        scale = min(120 / h, 180 / w)
        img = XLImage(io.BytesIO(it["full"]))
        img.width, img.height = int(w * scale), int(h * scale)
        ws.add_image(img, f"E{i}")
        ws.row_dimensions[i].height = img.height * 0.75 + 4
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main(page: ft.Page):
    page.title = "PhotoBase"
    page.theme_mode = ft.ThemeMode.SYSTEM
    page.padding = 0
    items: list[dict] = []
    state = {"q": "", "bytes": None}
    mobile = page.platform in (ft.PagePlatform.ANDROID, ft.PagePlatform.IOS)

    def snack(msg):
        page.open(ft.SnackBar(ft.Text(msg)))

    def persist():
        try:
            tmp = BASE.with_suffix(".tmp")
            tmp.write_bytes(write_xlsx(items))
            os.replace(tmp, BASE)
        except Exception as ex:
            snack(f"Sauvegarde locale impossible : {ex}")

    grid = ft.GridView(expand=True, max_extent=170, spacing=4, run_spacing=4, padding=6)

    def refresh():
        q = state["q"].strip().lower()
        grid.controls.clear()
        for it in reversed(items):
            if q and q not in " ".join(it[f] for f in FIELDS).lower():
                continue
            grid.controls.append(ft.Container(
                ft.Image(src_base64=base64.b64encode(it["thumb"]).decode(), fit=ft.ImageFit.COVER),
                border_radius=8, clip_behavior=ft.ClipBehavior.ANTI_ALIAS,
                on_click=lambda e, row=it: open_detail(row)))
        count.value = f"{len(items)} photo(s)"
        page.update()

    # --- Import images ---
    def on_images(e: ft.FilePickerResultEvent):
        n = 0
        for f in e.files or []:
            try:
                items.append(new_item(nom=Path(f.name).stem[:80], raw=Path(f.path).read_bytes()))
                n += 1
            except Exception as ex:
                snack(f"{f.name} : {ex}")
        if n:
            persist()
            snack(f"{n} photo(s) ajoutée(s)")
        refresh()

    # --- Import Excel ---
    def on_xlsx(e: ft.FilePickerResultEvent):
        if not e.files:
            return
        try:
            new = read_xlsx(e.files[0].path)
        except Exception as ex:
            return snack(f"Fichier illisible : {ex}")
        if not new:
            return snack("Aucune image trouvée. Utilise Insertion > Images > Placer sur les cellules.")
        items.extend(new)
        persist()
        snack(f"{len(new)} image(s) importée(s)")
        refresh()

    # --- Export Excel ---
    def on_saved(e: ft.FilePickerResultEvent):
        if not mobile and e.path:
            Path(e.path).write_bytes(state["bytes"])
        if e.path or mobile:
            snack("Export terminé")

    pick_img = ft.FilePicker(on_result=on_images)
    pick_xlsx = ft.FilePicker(on_result=on_xlsx)
    saver = ft.FilePicker(on_result=on_saved)
    page.overlay.extend([pick_img, pick_xlsx, saver])

    def do_import(_):
        pick_xlsx.pick_files(allowed_extensions=["xlsx"], file_type=ft.FilePickerFileType.CUSTOM)

    def do_export(_):
        if not items:
            return snack("Rien à exporter")
        state["bytes"] = write_xlsx(items)
        if mobile:
            saver.save_file(file_name="photos_base.xlsx", src_bytes=state["bytes"])
        else:
            saver.save_file(file_name="photos_base.xlsx", allowed_extensions=["xlsx"])

    def do_clear(_):
        def yes(_):
            items.clear()
            persist()
            page.close(dlg)
            refresh()

        dlg = ft.AlertDialog(
            title=ft.Text("Vider toute la base ?"),
            actions=[ft.TextButton("Annuler", on_click=lambda _: page.close(dlg)),
                     ft.TextButton("Vider", on_click=yes)])
        page.open(dlg)

    # --- Détail ---
    def open_detail(it):
        fields = {f: ft.TextField(label=f.capitalize(), value=it[f]) for f in FIELDS}

        def save(_):
            for f, t in fields.items():
                it[f] = t.value[:500]
            persist()
            page.close(dlg)
            refresh()

        def delete(_):
            items.remove(it)
            persist()
            page.close(dlg)
            refresh()

        dlg = ft.AlertDialog(
            scrollable=True,
            content=ft.Column(
                [ft.Image(src_base64=base64.b64encode(it["full"]).decode(), fit=ft.ImageFit.CONTAIN, height=320)]
                + list(fields.values()), tight=True, spacing=10),
            actions=[ft.TextButton("Supprimer", on_click=delete),
                     ft.TextButton("Fermer", on_click=lambda _: page.close(dlg)),
                     ft.FilledButton("Enregistrer", on_click=save)])
        page.open(dlg)

    # --- Écran principal ---
    count = ft.Text(size=12)

    def on_search(e):
        state["q"] = e.control.value
        refresh()

    page.appbar = ft.AppBar(
        title=ft.Text("PhotoBase"),
        actions=[ft.PopupMenuButton(items=[
            ft.PopupMenuItem(text="Importer un Excel", icon=ft.Icons.UPLOAD_FILE, on_click=do_import),
            ft.PopupMenuItem(text="Exporter en Excel", icon=ft.Icons.DOWNLOAD, on_click=do_export),
            ft.PopupMenuItem(text="Vider la base", icon=ft.Icons.DELETE_FOREVER, on_click=do_clear),
        ])])
    page.floating_action_button = ft.FloatingActionButton(
        icon=ft.Icons.ADD_A_PHOTO,
        on_click=lambda _: pick_img.pick_files(allow_multiple=True, file_type=ft.FilePickerFileType.IMAGE))
    search = ft.TextField(hint_text="Rechercher", prefix_icon=ft.Icons.SEARCH, dense=True,
                          border_radius=20, on_change=on_search)
    page.add(ft.SafeArea(ft.Column([ft.Container(ft.Column([search, count], spacing=4),
                                                 padding=ft.padding.symmetric(10, 12)), grid],
                                   expand=True), expand=True))

    if BASE.exists():
        try:
            items.extend(read_xlsx(BASE))
        except Exception:
            pass
    refresh()


ft.app(target=main)
